"""Tests du module d'extraction des pieces jointes (plan P2bis) — sans reseau."""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "Streamlit Cloud"))

from attachments import AttachmentError, extract_attachment, format_for_prompt  # noqa: E402


def _docx_bytes(paragraphs):
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Devis"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _blank_pdf_bytes():
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def test_docx_extraction():
    att = extract_attachment("devis.docx", _docx_bytes(["Devis toiture", "Montant : 12 000 EUR HT"]))
    assert att["kind"] == "docx" and not att["truncated"]
    assert "12 000 EUR HT" in att["text"]


def test_xlsx_extraction_tableaux():
    att = extract_attachment("budget.xlsx", _xlsx_bytes([["Poste", "Montant"], ["Toiture", 12000]]))
    assert "Feuille : Devis" in att["text"]
    assert "Toiture | 12000" in att["text"]


def test_csv_extraction_accents():
    att = extract_attachment("lots.csv", "lot;coproprietaire\n12;Müller".encode("cp1252"))
    assert "Müller" in att["text"]


def test_pdf_scanne_refuse_avec_message_honnete():
    with pytest.raises(AttachmentError, match="scann"):
        extract_attachment("scan.pdf", _blank_pdf_bytes())


def test_format_non_supporte_et_doc_legacy():
    with pytest.raises(AttachmentError, match="non supporte"):
        extract_attachment("photo.png", b"x")
    with pytest.raises(AttachmentError, match=".docx"):
        extract_attachment("vieux.doc", b"x")


def test_troncature_et_bloc_prompt():
    att = extract_attachment("gros.docx", _docx_bytes(["ligne " + "x" * 90] * 600))
    assert att["truncated"]
    bloc = format_for_prompt(att)
    assert bloc.startswith("[Document joint par l'utilisateur : gros.docx (tronque)")
    assert "PAS un document de la base documentaire" in bloc
    assert bloc.rstrip().endswith("<<<FIN PIECE JOINTE>>>")
